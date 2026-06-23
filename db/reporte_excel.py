from io import BytesIO
from datetime import datetime, timezone, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path
from openpyxl.drawing.image import Image as XLImage

from db.samee100_db import get_conn
from db.kpi_solar import reporte_kpi_energetico, rango_real_datos

BASE_DIR = Path(__file__).resolve().parent.parent
LOGO_PATH = BASE_DIR / "static" / "img" / "pyp_logo.jpg"

UNIT_IDS_KPI = {
    "PTotal": 61,
    "EPImp": 100,
    "EPExp": 104,
    "P1": 58,
    "P2": 59,
    "P3": 60,
    "EP1Imp": 97,
    "EP2Imp": 98,
    "EP3Imp": 99,
    "EP1Exp": 101,
    "EP2Exp": 102,
    "EP3Exp": 103,
}


def convertir_utc_a_colombia(timestamp_utc):
    try:
        if timestamp_utc is None:
            return ""

        if str(timestamp_utc).isdigit():
            dt_utc = datetime.fromtimestamp(
                int(timestamp_utc),
                tz=timezone.utc
            )
        else:
            texto = str(timestamp_utc).replace("Z", "+00:00")
            dt_utc = datetime.fromisoformat(texto)

            if dt_utc.tzinfo is None:
                dt_utc = dt_utc.replace(tzinfo=timezone.utc)

        dt_col = dt_utc.astimezone(timezone(timedelta(hours=-5)))
        return dt_col.strftime("%Y-%m-%d %H:%M:%S")

    except Exception:
        return ""


def aplicar_estilo_titulo(ws, celda, texto):
    ws[celda] = texto
    ws[celda].font = Font(bold=True, size=16, color="FFFFFF")
    ws[celda].fill = PatternFill("solid", fgColor="0F3B63")
    ws[celda].alignment = Alignment(horizontal="center")


def aplicar_header(ws, fila):
    fill = PatternFill("solid", fgColor="0F3B63")
    font = Font(bold=True, color="FFFFFF")
    border = Border(
        bottom=Side(style="thin", color="D9E2EC")
    )

    for cell in ws[fila]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border


def ajustar_columnas(ws):
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)

        for cell in col:
            value = str(cell.value) if cell.value is not None else ""
            if len(value) > max_length:
                max_length = len(value)

        ws.column_dimensions[col_letter].width = min(max_length + 3, 45)


def obtener_variable(unit_id):
    conn = get_conn()
    cur = conn.cursor()

    cur.execute("""
        SELECT unit_id, name AS variable, simbol AS simbolo
        FROM unidades
        WHERE unit_id = ?
    """, (unit_id,))

    row = cur.fetchone()
    conn.close()

    if not row:
        return {
            "unit_id": unit_id,
            "variable": f"Variable {unit_id}",
            "simbolo": ""
        }

    return dict(row)


def obtener_delta_energia(unit_id, inicio, fin):
    conn = get_conn()
    cur = conn.cursor()

    cur.execute("""
        SELECT valor, timestamp_utc
        FROM mediciones_detalle
        WHERE unit_id = ?
          AND timestamp_utc >= ?
          AND timestamp_utc <= ?
        ORDER BY timestamp_utc ASC, id ASC
        LIMIT 1
    """, (unit_id, inicio, fin))
    inicial = cur.fetchone()

    cur.execute("""
        SELECT valor, timestamp_utc
        FROM mediciones_detalle
        WHERE unit_id = ?
          AND timestamp_utc >= ?
          AND timestamp_utc <= ?
        ORDER BY timestamp_utc DESC, id DESC
        LIMIT 1
    """, (unit_id, inicio, fin))
    final = cur.fetchone()

    conn.close()

    if not inicial or not final:
        return {
            "inicial": None,
            "final": None,
            "delta": 0.0,
            "timestamp_inicial": None,
            "timestamp_final": None
        }

    delta = float(final["valor"]) - float(inicial["valor"])

    if delta < 0:
        delta = 0.0

    return {
        "inicial": round(float(inicial["valor"]), 3),
        "final": round(float(final["valor"]), 3),
        "delta": round(delta, 3),
        "timestamp_inicial": inicial["timestamp_utc"],
        "timestamp_final": final["timestamp_utc"]
    }


def obtener_series_variables(unit_ids, inicio, fin, limite=50000):
    conn = get_conn()
    cur = conn.cursor()

    placeholders = ",".join(["?"] * len(unit_ids))

    sql = f"""
        SELECT
            md.timestamp_utc,
            md.device_id,
            md.unit_id,
            u.name AS variable,
            u.simbol AS simbolo,
            md.valor
        FROM mediciones_detalle md
        LEFT JOIN unidades u
            ON md.unit_id = u.unit_id
        WHERE md.unit_id IN ({placeholders})
          AND md.timestamp_utc >= ?
          AND md.timestamp_utc <= ?
        ORDER BY md.timestamp_utc ASC, md.unit_id ASC
        LIMIT ?
    """

    params = list(unit_ids) + [inicio, fin, limite]

    cur.execute(sql, params)
    rows = [dict(r) for r in cur.fetchall()]

    conn.close()
    return rows


def obtener_todas_las_variables_con_datos():
    conn = get_conn()
    cur = conn.cursor()

    cur.execute("""
        SELECT DISTINCT unit_id
        FROM mediciones_detalle
        ORDER BY unit_id ASC
    """)

    ids = [int(r["unit_id"]) for r in cur.fetchall()]

    conn.close()
    return ids


def crear_reporte_excel(inicio, fin, variables_param="61,104,100"):
    wb = Workbook()

    # Eliminar hoja por defecto
    ws_default = wb.active
    wb.remove(ws_default)

    data_kpi = reporte_kpi_energetico(inicio, fin)
    rango = rango_real_datos(inicio, fin)

    inicio_real = rango.get("inicio_real")
    fin_real = rango.get("fin_real")

    # =========================================================
    # HOJA 1: RESUMEN KPI
    # =========================================================
    ws = wb.create_sheet("Resumen_KPI")

    # Encabezado ejecutivo
    insertar_logo(ws, "A1")

    ws.merge_cells("B1:F1")
    ws["B1"] = "REPORTE ENERGÉTICO SAMEE100"
    ws["B1"].font = Font(bold=True, size=18, color="0F3B63")
    ws["B1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("B2:F2")
    ws["B2"] = "Sistema de Monitoreo Energético - PYP Tecnología Electrónica SAS"
    ws["B2"].font = Font(bold=True, size=11, color="1F2937")
    ws["B2"].alignment = Alignment(horizontal="center")

    ws.row_dimensions[1].height = 42
    ws.row_dimensions[2].height = 24

    ws.append([])
    ws.append(["Periodo consulta UTC", inicio, fin])
    ws.append([
        "Periodo real Colombia",
        convertir_utc_a_colombia(inicio_real),
        convertir_utc_a_colombia(fin_real)
    ])
    ws.append([
        "Fecha generación reporte Colombia",
        convertir_utc_a_colombia(str(int(datetime.now(timezone.utc).timestamp()))),
        ""
    ])

    ws.append([])
    ws.append(["Indicador", "Valor", "Unidad", "Descripción"])
    aplicar_header(ws, 8)

    energia = data_kpi["energia"]
    potencia = data_kpi["potencia"]
    impacto = data_kpi["impacto"]

    filas_kpi = [
        ["Energía exportada / generada", energia["exportada_kwh"], "kWh", "Energía activa inversa total EPExp"],
        ["Energía importada de red", energia["importada_kwh"], "kWh", "Energía activa positiva total EPImp"],
        ["Balance neto", energia["balance_neto_kwh"], "kWh", "Exportada - Importada"],
        ["Ahorro estimado", impacto["ahorro_cop"], "COP", "Cálculo con tarifa configurada"],
        ["CO₂ evitado", impacto["co2_evitado_kg"], "kg", "Factor de emisión configurado"],
        ["Potencia actual", potencia["actual_kw"], "kW", "Último valor PTotal"],
        ["Potencia máxima", potencia["maxima_kw"], "kW", "Máximo PTotal en el periodo"],
        ["Potencia mínima", potencia["minima_kw"], "kW", "Mínimo PTotal en el periodo"],
        ["Potencia promedio", potencia["promedio_kw"], "kW", "Promedio PTotal en el periodo"],
        ["Muestras potencia", potencia["muestras"], "registros", "Cantidad de muestras PTotal"]
    ]

    for fila in filas_kpi:
        ws.append(fila)

    ajustar_columnas(ws)

    # =========================================================
    # HOJA 2: ENERGIA FASES
    # =========================================================
    ws2 = wb.create_sheet("Energia_Fases")

    ws2.merge_cells("A1:H1")
    aplicar_estilo_titulo(ws2, "A1", "ENERGÍA POR FASES Y TOTALES")

    ws2.append([])
    ws2.append([
        "Unit ID",
        "Variable",
        "Tipo",
        "Inicial kWh",
        "Final kWh",
        "Delta kWh",
        "Hora inicial Colombia",
        "Hora final Colombia"
    ])
    aplicar_header(ws2, 3)

    energia_units = [
        (97, "Importada L1"),
        (98, "Importada L2"),
        (99, "Importada L3"),
        (100, "Importada Total"),
        (101, "Exportada L1"),
        (102, "Exportada L2"),
        (103, "Exportada L3"),
        (104, "Exportada Total")
    ]

    for unit_id, tipo in energia_units:
        info_var = obtener_variable(unit_id)
        delta = obtener_delta_energia(unit_id, inicio, fin)

        ws2.append([
            unit_id,
            info_var["variable"],
            tipo,
            delta["inicial"],
            delta["final"],
            delta["delta"],
            convertir_utc_a_colombia(delta["timestamp_inicial"]),
            convertir_utc_a_colombia(delta["timestamp_final"])
        ])

    ajustar_columnas(ws2)

    # =========================================================
    # HOJA 3: POTENCIA
    # =========================================================
    ws3 = wb.create_sheet("Potencia")

    ws3.merge_cells("A1:F1")
    aplicar_estilo_titulo(ws3, "A1", "ANÁLISIS DE POTENCIA")

    ws3.append([])
    ws3.append(["Indicador", "Valor", "Unidad"])
    aplicar_header(ws3, 3)

    ws3.append(["Potencia actual", potencia["actual_kw"], "kW"])
    ws3.append(["Potencia máxima", potencia["maxima_kw"], "kW"])
    ws3.append(["Potencia mínima", potencia["minima_kw"], "kW"])
    ws3.append(["Potencia promedio", potencia["promedio_kw"], "kW"])
    ws3.append(["Muestras", potencia["muestras"], "registros"])

    ajustar_columnas(ws3)

    # =========================================================
    # HOJA 4: VARIABLES
    # =========================================================
    ws4 = wb.create_sheet("Variables")

    ws4.merge_cells("A1:G1")
    aplicar_estilo_titulo(ws4, "A1", "HISTÓRICO DE VARIABLES SELECCIONADAS")

    ws4.append([])

    if variables_param == "all":
        unit_ids = obtener_todas_las_variables_con_datos()
    else:
        unit_ids = [
            int(x.strip())
            for x in variables_param.split(",")
            if x.strip().isdigit()
        ]

    if not unit_ids:
        unit_ids = [61, 100, 104]

    ws4.append([
        "Timestamp UTC",
        "Hora Colombia",
        "Device ID",
        "Unit ID",
        "Variable",
        "Unidad",
        "Valor"
    ])
    aplicar_header(ws4, 3)

    rows = obtener_series_variables(unit_ids, inicio, fin)

    for r in rows:
        ws4.append([
            r["timestamp_utc"],
            convertir_utc_a_colombia(r["timestamp_utc"]),
            r["device_id"],
            r["unit_id"],
            r["variable"],
            r["simbolo"],
            r["valor"]
        ])

    ajustar_columnas(ws4)

    # =========================================================
    # HOJA 5: METADATOS
    # =========================================================
    ws5 = wb.create_sheet("Metadatos")

    ws5.merge_cells("A1:D1")
    aplicar_estilo_titulo(ws5, "A1", "METADATOS DEL REPORTE")

    ws5.append([])
    ws5.append(["Campo", "Valor"])
    aplicar_header(ws5, 3)

    ws5.append(["Sistema", "SAMEE100"])
    ws5.append(["Empresa", "PYP Tecnología Electrónica SAS"])
    ws5.append(["Inicio consulta UTC", inicio])
    ws5.append(["Fin consulta UTC", fin])
    ws5.append(["Inicio real Colombia", convertir_utc_a_colombia(inicio_real)])
    ws5.append(["Fin real Colombia", convertir_utc_a_colombia(fin_real)])
    ws5.append(["Variables incluidas", variables_param])
    ws5.append(["Hojas", "Resumen_KPI, Energia_Fases, Potencia, Variables, Metadatos"])

    ajustar_columnas(ws5)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output

def insertar_logo(ws, celda="A1"):
    try:
        if LOGO_PATH.exists():
            img = XLImage(str(LOGO_PATH))
            img.width = 110
            img.height = 55
            ws.add_image(img, celda)
    except Exception as e:
        print(f"[REPORTE] No se pudo insertar logo: {e}")